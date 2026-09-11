"""Build an editable PowerPoint deck from the generated Beamer source.

``build_slides.py`` writes ``slides/pedrad_ai_slides.tex``; this script reads
that file (after all ``@@KEY@@`` substitutions) and re-creates each frame as a
native PowerPoint slide:

- every ``\\includegraphics`` becomes a picture shape (drag, resize, crop),
- itemize / enumerate / paragraphs become text boxes,
- ``tabular`` becomes a real PowerPoint table,
- ``columns`` are laid out side by side,
- a small "n / N" page counter sits bottom-right, matching the Beamer footline.

It understands only the LaTeX subset the template uses (see ``build_slides.py``),
so add a case here if a new construct appears in the template.

Usage::

    python scripts/build_pptx.py                  # slides/pedrad_ai_slides.pptx
    python scripts/build_pptx.py --tex path.tex --out deck.pptx
"""
from __future__ import annotations

import argparse
import math
import re
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image as PILImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SLIDES_DIR = ROOT / "slides"

# ---------------------------------------------------------------- geometry ---
SLIDE_W, SLIDE_H = 13.333, 7.5          # inches, 16:9
MARGIN_X = 0.45
TITLE_BAND_H = 0.78
BODY_TOP = TITLE_BAND_H + 0.18
BODY_BOTTOM = SLIDE_H - 0.40
BODY_W = SLIDE_W - 2 * MARGIN_X
BODY_H = BODY_BOTTOM - BODY_TOP
GAP = 0.08                              # vertical gap between blocks
COL_GAP = 0.2

STRUCTURE = RGBColor(0x33, 0x33, 0xB3)  # beamer "structure" blue (whale)
TEXT = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
BLOCK_FILL = RGBColor(0xE8, 0xEA, 0xF6)
FONT = "Calibri"

SIZE_PT = {"normal": 16, "large": 20, "small": 14, "footnotesize": 12, "scriptsize": 11, "tiny": 9}
SIZE_WORDS = "tiny|scriptsize|footnotesize|small|normalsize|large"


# ------------------------------------------------------------- data model ---
Run = tuple[str, bool, bool]             # text, bold, italic
Paragraph = list[Run]


@dataclass
class Text:
    paragraphs: list[Paragraph]
    size: str = "normal"
    align: str = "left"
    bullet: str | None = None           # None | "bullet" | "number"
    boxed: bool = False


@dataclass
class Image:
    path: Path
    height_frac: float = 0.8
    align: str = "center"


@dataclass
class Table:
    rows: list[list[Paragraph]]
    colspec: list[float | None]         # fraction of width, or None = auto
    size: str = "normal"


@dataclass
class Columns:
    cols: list[tuple[float, list]] = field(default_factory=list)   # (width frac, blocks)


# --------------------------------------------------------------- tex utils ---
def find_matching(s: str, i: int) -> int:
    """Index of the '}' matching the '{' at s[i]."""
    assert s[i] == "{", s[i:i + 20]
    depth = 0
    j = i
    while j < len(s):
        c = s[j]
        if c == "\\":
            j += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth == 0:
                return j
        j += 1
    raise ValueError("unbalanced braces")


def find_env_end(s: str, env: str, start: int) -> tuple[int, int]:
    """(start, end) of the matching ``\\end{env}`` for a body beginning at ``start``."""
    depth = 1
    pat = re.compile(r"\\(begin|end)\{" + re.escape(env) + r"\}")
    for m in pat.finditer(s, start):
        depth += 1 if m.group(1) == "begin" else -1
        if depth == 0:
            return m.start(), m.end()
    raise ValueError(f"no \\end{{{env}}}")


MATH = {
    r"\geq": "\u2265", r"\leq": "\u2264", r"\cap": "\u2229", r"\cup": "\u222a", r"\approx": "\u2248",
    r"\rightarrow": "\u2192", r"\leftarrow": "\u2190", r"\times": "\u00d7", r"\pm": "\u00b1", r"\sim": "~",
    r"\infty": "\u221e", r"\star": "\u2605", "<": "<", ">": ">", "+": "+", "-": "\u2212", "=": "=",
}


def _math(m: re.Match) -> str:
    inner = m.group(1).strip()
    out = inner
    for k in sorted(MATH, key=len, reverse=True):
        out = out.replace(k, MATH[k])
    return out.replace("\\", "").replace("{", "").replace("}", "")


def clean_text(s: str) -> str:
    """Plain-text rendering of a LaTeX fragment with no formatting commands left."""
    s = re.sub(r"\$([^$]*)\$", _math, s)
    s = re.sub(r"\\(ldots|dots|cdots)(\{\})?", "\u2026", s)
    s = re.sub(r"\\(hspace|vspace)\*?\{[^}]*\}", " ", s)
    # Layout declarations carry no content; without this the argument of
    # \renewcommand{\arraystretch}{0.92} lands on the slide as a stray "0.92".
    s = re.sub(r"\\renewcommand\{[^}]*\}\{[^}]*\}", "", s)
    s = re.sub(r"\\(hfill|centering|noindent|small|tiny|scriptsize|footnotesize|normalsize|large)\b", "", s)
    s = s.replace("\\\\", " ")
    s = s.replace("---", "\u2014").replace("--", "\u2013")
    s = s.replace("``", "\u201c").replace("''", "\u201d").replace("`", "\u2018")
    s = s.replace("\\&", "&").replace("\\%", "%").replace("\\_", "_").replace("\\#", "#").replace("\\$", "$")
    s = s.replace("\\,", "\u2009").replace("~", "\u00a0")
    s = re.sub(r"\\[A-Za-z]+\*?", "", s)      # any leftover command name
    s = s.replace("{", "").replace("}", "")
    s = re.sub(r"\s+", " ", s)
    return s


FMT_RE = re.compile(r"\\(textbf|textit|emph|textsc)\{")


def tex_to_runs(s: str, bold: bool = False, italic: bool = False) -> Paragraph:
    runs: Paragraph = []
    pos = 0
    while True:
        m = FMT_RE.search(s, pos)
        if not m:
            runs.append((clean_text(s[pos:]), bold, italic))
            break
        runs.append((clean_text(s[pos:m.start()]), bold, italic))
        close = find_matching(s, m.end() - 1)
        inner = s[m.end():close]
        cmd = m.group(1)
        runs.extend(tex_to_runs(inner, bold or cmd == "textbf", italic or cmd in ("textit", "emph")))
        pos = close + 1
    # merge / drop empties, trim outer whitespace
    out: Paragraph = []
    for t, b, i in runs:
        if not t:
            continue
        if out and out[-1][1:] == (b, i):
            out[-1] = (out[-1][0] + t, b, i)
        else:
            out.append((t, b, i))
    if out:
        out[0] = (out[0][0].lstrip(), *out[0][1:])
        out[-1] = (out[-1][0].rstrip(), *out[-1][1:])
    return [r for r in out if r[0]]


def split_paragraphs(s: str) -> list[Paragraph]:
    parts = re.split(r"\\\\(?:\[[^\]]*\])?|\n\s*\n", s)
    paras = [tex_to_runs(p) for p in parts]
    return [p for p in paras if p]


# ------------------------------------------------------------------ parser ---
SPECIAL = re.compile(
    r"\\begin\{(?P<env>\w+)\}"
    r"|\\includegraphics"
    r"|\{\\(?P<gsize>" + SIZE_WORDS + r")\b"
    r"|\\(?P<size>" + SIZE_WORDS + r")\b"
    r"|\\centering\b"
    r"|\\vspace\*?\{[^}]*\}"
)


def _frac(s: str, default: float) -> float:
    m = re.search(r"([0-9.]+)\s*\\(textwidth|textheight|linewidth|columnwidth)", s)
    return float(m.group(1)) if m else default


def parse_blocks(src: str, size: str = "normal", align: str = "left", fig_dir: Path | None = None) -> list:
    blocks: list = []
    pos = 0

    def flush(text: str) -> None:
        paras = split_paragraphs(text)
        if paras:
            blocks.append(Text(paras, size=size, align=align))

    while True:
        m = SPECIAL.search(src, pos)
        if not m:
            flush(src[pos:])
            break
        flush(src[pos:m.start()])
        if m.group("env"):
            env = m.group("env")
            body_start = m.end()
            # optional [..] and {..} arguments
            args: list[str] = []
            opt = ""
            while True:
                if src.startswith("[", body_start):
                    j = src.index("]", body_start)
                    opt = src[body_start + 1:j]
                    body_start = j + 1
                elif src.startswith("{", body_start) and env in ("tabular", "column", "block", "minipage"):
                    j = find_matching(src, body_start)
                    args.append(src[body_start + 1:j])
                    body_start = j + 1
                    if env == "tabular" or env == "column" or env == "block" or len(args) >= 2:
                        break
                else:
                    break
            end_start, end_end = find_env_end(src, env, body_start)
            inner = src[body_start:end_start]
            blocks.extend(parse_env(env, args, opt, inner, size, align, fig_dir))
            pos = end_end
        elif m.group(0) == "\\includegraphics":
            j = m.end()
            opts = ""
            if src.startswith("[", j):
                k = src.index("]", j)
                opts = src[j + 1:k]
                j = k + 1
            k = find_matching(src, j)
            fname = src[j + 1:k].strip()
            path = (fig_dir / fname) if fig_dir else Path(fname)
            blocks.append(Image(path, height_frac=_frac(opts, 0.8), align=align if align != "left" else "center"))
            pos = k + 1
        elif m.group("gsize"):
            close = find_matching(src, m.start())
            inner = src[m.end():close]
            gsize = m.group("gsize").replace("normalsize", "normal")
            blocks.extend(parse_blocks(inner, gsize, align, fig_dir))
            pos = close + 1
        elif m.group("size"):
            size = m.group("size").replace("normalsize", "normal")
            pos = m.end()
        elif m.group(0).startswith("\\centering"):
            align = "center"
            pos = m.end()
        else:                                   # \vspace: ignore
            pos = m.end()
    return blocks


def parse_env(env: str, args: list[str], opt: str, inner: str, size: str, align: str, fig_dir) -> list:
    if env in ("itemize", "enumerate"):
        items = [it for it in re.split(r"\\item\b", inner) if it.strip()]
        paras = [tex_to_runs(it) for it in items]
        return [Text([p for p in paras if p], size=size, align="left",
                     bullet="number" if env == "enumerate" else "bullet")]
    if env == "center":
        return parse_blocks(inner, size, "center", fig_dir)
    if env == "columns":
        cols = []
        col_re = re.compile(r"\\begin\{column\}(\[[^\]]*\])?\{")
        pos = 0
        while True:
            m = col_re.search(inner, pos)
            if not m:
                break
            j = find_matching(inner, m.end() - 1)
            width = _frac(inner[m.end():j], 0.5)
            es, ee = find_env_end(inner, "column", j + 1)
            cols.append((width, parse_blocks(inner[j + 1:es], size, align, fig_dir)))
            pos = ee
        return [Columns(cols)]
    if env == "column":                         # bare column outside columns: treat as flow
        return parse_blocks(inner, size, align, fig_dir)
    if env == "block":
        out = parse_blocks(inner, size, align, fig_dir)
        title = args[0].strip() if args else ""
        if title:
            out.insert(0, Text([tex_to_runs(r"\textbf{" + title + "}")], size=size))
        for b in out:
            if isinstance(b, Text):
                b.boxed = True
        return out
    if env == "tabular":
        return [parse_table(args[0] if args else "", inner, size)]
    if env == "minipage":
        return parse_blocks(inner, size, align, fig_dir)
    # unknown environment: keep its text
    return parse_blocks(inner, size, align, fig_dir)


def parse_table(colspec: str, inner: str, size: str) -> Table:
    spec: list[float | None] = []
    # Drop the >{...} / <{...} / @{...} decorations first: they are not columns,
    # and the letters inside them ("raggedright", "arraybackslash") would each
    # match the bare-column alternative and invent a column per letter.
    colspec = re.sub(r"[><@]\{[^{}]*\}", "", colspec)
    for m in re.finditer(r"p\{([^}]*)\}|([lcr])", colspec):
        spec.append(_frac(m.group(1), 0.2) if m.group(1) else None)
    body = re.sub(r"\\(hline|toprule|midrule|bottomrule)", "", inner)
    rows: list[list[Paragraph]] = []
    for raw in re.split(r"\\\\(?:\[[^\]]*\])?", body):
        if not raw.strip():
            continue
        cells = re.split(r"(?<!\\)&", raw)
        rows.append([tex_to_runs(c) for c in cells])
    ncol = max(len(spec), max(len(r) for r in rows)) if rows else len(spec)
    spec += [None] * (ncol - len(spec))
    for r in rows:
        r += [[] for _ in range(ncol - len(r))]
    return Table(rows, spec, size=size)


# ----------------------------------------------------------- deck parsing ---
@dataclass
class Frame:
    title: str
    blocks: list
    title_slide: bool = False


def parse_deck(tex: str, fig_dir: Path) -> tuple[str, str, str, str, list[Frame]]:
    def _cmd(name: str) -> str:
        m = re.search(r"\\" + name + r"(?:\[[^\]]*\])?\{", tex)
        if not m:
            return ""
        return clean_text(tex[m.end():find_matching(tex, m.end() - 1)])

    title, subtitle = _cmd("title"), _cmd("subtitle")
    author, date = _cmd("author"), _cmd("date")
    body = tex[tex.index(r"\begin{document}"):]
    frames: list[Frame] = []
    frame_re = re.compile(r"\\frame\{\\titlepage\}|\\begin\{frame\}(?:\[[^\]]*\])?")
    pos = 0
    while True:
        m = frame_re.search(body, pos)
        if not m:
            break
        if m.group(0).startswith(r"\frame{"):
            frames.append(Frame(title, [], title_slide=True))
            pos = m.end()
            continue
        j = m.end()
        ftitle = ""
        if body.startswith("{", j):
            k = find_matching(body, j)
            ftitle = clean_text(body[j + 1:k])
            j = k + 1
        es, ee = find_env_end(body, "frame", j)
        frames.append(Frame(ftitle, parse_blocks(body[j:es], fig_dir=fig_dir)))
        pos = ee
    return title, subtitle, author, date, frames


# ------------------------------------------------------------------ layout ---
def _pt(size: str, scale: float) -> float:
    return SIZE_PT.get(size, 16) * scale


def _lines(text_len: int, pt: float, width_in: float) -> int:
    char_w = 0.5 * pt / 72                     # average glyph width, inches
    per_line = max(1, int(width_in / char_w))
    return max(1, math.ceil(text_len / per_line))


def text_height(block: Text, width_in: float, scale: float) -> float:
    pt = _pt(block.size, scale)
    indent = 0.3 if block.bullet else 0.0
    h = 0.0
    for para in block.paragraphs:
        n = sum(len(t) for t, _, _ in para)
        h += _lines(n, pt, width_in - indent - 0.2) * pt * 1.2 / 72 + pt * 0.35 / 72
    return h + 0.1


def image_box(block: Image, width_in: float, scale: float) -> tuple[float, float]:
    try:
        with PILImage.open(block.path) as im:
            ar = im.width / im.height
    except Exception:
        ar = 1.6
    h = block.height_frac * BODY_H * scale
    w = h * ar
    if w > width_in:
        w = width_in
        h = w / ar
    return w, h


def table_widths(block: Table, width_in: float) -> list[float]:
    auto = 0.75
    fixed = sum(auto for s in block.colspec if s is None)
    fracs = [s for s in block.colspec if s is not None]
    rem = max(width_in - fixed, 1.0)
    total = sum(fracs) or 1.0
    return [auto if s is None else rem * s / total for s in block.colspec]


def table_height(block: Table, widths: list[float], scale: float) -> tuple[list[float], float]:
    pt = _pt(block.size, scale)
    heights = []
    for row in block.rows:
        lines = 1
        for cell, w in zip(row, widths):
            n = sum(len(t) for t, _, _ in cell)
            lines = max(lines, _lines(n, pt, w - 0.15))
        heights.append(lines * pt * 1.2 / 72 + 0.09)
    return heights, sum(heights)


def layout(blocks: list, left: float, top: float, width: float, scale: float) -> tuple[list, float]:
    """Place blocks top-down; return (placements, total height)."""
    items: list = []
    y = top
    for b in blocks:
        if isinstance(b, Text):
            h = text_height(b, width, scale)
            items.append(("text", left, y, width, h, b))
            y += h + GAP
        elif isinstance(b, Image):
            w, h = image_box(b, width, scale)
            x = left + (width - w) / 2 if b.align == "center" else left
            items.append(("image", x, y, w, h, b))
            y += h + GAP
        elif isinstance(b, Table):
            widths = table_widths(b, width)
            rh, h = table_height(b, widths, scale)
            items.append(("table", left, y, sum(widths), h, (b, widths, rh)))
            y += h + GAP
        elif isinstance(b, Columns):
            fr = sum(f for f, _ in b.cols) or 1.0
            n = len(b.cols)
            gap = COL_GAP if n > 1 else 0.0
            avail = width - gap * (n - 1)
            x = left
            hmax = 0.0
            for f, sub in b.cols:
                cw = avail * (f / max(fr, 1.0)) if fr <= 1.0 else avail * f / fr
                sub_items, h = layout(sub, x, y, cw, scale)
                items.extend(sub_items)
                hmax = max(hmax, h)
                x += cw + gap
            y += hmax + GAP
    return items, y - top - GAP


def fit_layout(blocks: list) -> list:
    for scale in (1.0, 0.94, 0.88, 0.82, 0.76, 0.7, 0.64, 0.58):
        items, h = layout(blocks, MARGIN_X, BODY_TOP, BODY_W, scale)
        if h <= BODY_H:
            return items
    return items


# ---------------------------------------------------------------- rendering ---
def _set_bullet(paragraph, kind: str, pt: float) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    indent = Emu(Inches(0.28))
    pPr.set("marL", str(int(indent)))
    pPr.set("indent", str(-int(indent)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if kind == "number":
        el = pPr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"})
    else:
        el = pPr.makeelement(qn("a:buChar"), {"char": "\u2022"})
    pPr.append(el)


def add_runs(paragraph, runs: Paragraph, pt: float, color=TEXT) -> None:
    for text, bold, italic in runs:
        r = paragraph.add_run()
        r.text = text
        r.font.size = Pt(pt)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = FONT
        r.font.color.rgb = color


def add_textbox(slide, x, y, w, h, block: Text, scale: float) -> None:
    pt = _pt(block.size, scale)
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    if block.boxed:
        tb.fill.solid()
        tb.fill.fore_color.rgb = BLOCK_FILL
        tf.margin_left = tf.margin_right = Inches(0.12)
        tf.margin_top = tf.margin_bottom = Inches(0.08)
    for i, para in enumerate(block.paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER if block.align == "center" else PP_ALIGN.LEFT
        p.space_after = Pt(pt * 0.35)
        if block.bullet:
            _set_bullet(p, block.bullet, pt)
        add_runs(p, para, pt)


def add_image(slide, x, y, w, h, block: Image) -> None:
    if not block.path.exists():
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tb.text_frame.text = f"[missing figure: {block.path.name}]"
        return
    pic = slide.shapes.add_picture(str(block.path), Inches(x), Inches(y), height=Inches(h))
    pic.name = block.path.stem


def add_table(slide, x, y, w, h, block: Table, widths: list[float], row_h: list[float], scale: float) -> None:
    pt = _pt(block.size, scale)
    nrows, ncols = len(block.rows), len(widths)
    shape = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(sum(widths)), Inches(h))
    tbl = shape.table
    for j, cw in enumerate(widths):
        tbl.columns[j].width = Inches(cw)
    for i, (row, rh) in enumerate(zip(block.rows, row_h)):
        tbl.rows[i].height = Inches(rh)
        for j, cell_runs in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            p = cell.text_frame.paragraphs[0]
            cell.text_frame.word_wrap = True
            header = i == 0
            add_runs(p, [(t, b or header, it) for t, b, it in cell_runs], pt,
                     color=RGBColor(0xFF, 0xFF, 0xFF) if header else TEXT)
            if block.colspec[j] is None:
                p.alignment = PP_ALIGN.RIGHT


def add_title_band(slide, title: str) -> None:
    band = slide.shapes.add_shape(1, 0, 0, Inches(SLIDE_W), Inches(TITLE_BAND_H))   # 1 = rectangle
    band.fill.solid()
    band.fill.fore_color.rgb = STRUCTURE
    band.line.fill.background()
    band.name = "title band"
    tf = band.text_frame
    tf.margin_left = Inches(MARGIN_X)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title
    r.font.size = Pt(22)
    r.font.name = FONT
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def add_page_number(slide, n: int, total: int) -> None:
    tb = slide.shapes.add_textbox(Inches(SLIDE_W - 1.6), Inches(SLIDE_H - 0.38), Inches(1.3), Inches(0.3))
    tb.name = "page number"
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = f"{n} / {total}"
    r.font.size = Pt(10)
    r.font.name = FONT
    r.font.color.rgb = MUTED


def render_title_slide(slide, title: str, subtitle: str, author: str = "", date: str = "") -> None:
    box = slide.shapes.add_shape(1, Inches(1.2), Inches(2.4), Inches(SLIDE_W - 2.4), Inches(1.3))
    box.fill.solid()
    box.fill.fore_color.rgb = STRUCTURE
    box.line.fill.background()
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.size = Pt(34)
    r.font.name = FONT
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    tb = slide.shapes.add_textbox(Inches(1.2), Inches(3.9), Inches(SLIDE_W - 2.4), Inches(1.0))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = subtitle
    r.font.size = Pt(18)
    r.font.name = FONT
    r.font.color.rgb = TEXT
    y = 4.9
    for line, size, color in ((author, 18, TEXT), (date, 14, MUTED)):
        if not line:
            continue
        tb = slide.shapes.add_textbox(Inches(1.2), Inches(y), Inches(SLIDE_W - 2.4), Inches(0.5))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.color.rgb = color
        y += 0.55


def build(tex_path: Path, out_path: Path) -> Path:
    tex = tex_path.read_text(encoding="utf-8")
    gp = re.search(r"\\graphicspath\{\{([^}]*)\}\}", tex)
    fig_dir = (tex_path.parent / gp.group(1)).resolve() if gp else tex_path.parent
    title, subtitle, author, date, frames = parse_deck(tex, fig_dir)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]
    total = len(frames)
    for n, fr in enumerate(frames, start=1):
        slide = prs.slides.add_slide(blank)
        if fr.title_slide:
            render_title_slide(slide, title, subtitle, author, date)
        else:
            add_title_band(slide, fr.title)
            scale_used = 1.0
            items = fit_layout(fr.blocks)
            # recover the scale fit_layout settled on (text size must match its height estimate)
            for scale in (1.0, 0.94, 0.88, 0.82, 0.76, 0.7, 0.64, 0.58):
                _, h = layout(fr.blocks, MARGIN_X, BODY_TOP, BODY_W, scale)
                if h <= BODY_H:
                    scale_used = scale
                    break
            for kind, x, y, w, h, payload in items:
                if kind == "text":
                    add_textbox(slide, x, y, w, h, payload, scale_used)
                elif kind == "image":
                    add_image(slide, x, y, w, h, payload)
                elif kind == "table":
                    tbl, widths, rh = payload
                    add_table(slide, x, y, w, h, tbl, widths, rh, scale_used)
        add_page_number(slide, n, total)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tex", type=Path, default=SLIDES_DIR / "pedrad_ai_slides.tex")
    ap.add_argument("--out", type=Path, default=None)
    args = ap.parse_args()
    out = args.out or args.tex.with_suffix(".pptx")
    build(args.tex, out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
